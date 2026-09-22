"""Выгрузка в Excel: лист «Цены» (текущий срез) и «Изменения» (история за 30 дней)."""
from __future__ import annotations

from datetime import datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from ..money import excel_format
from ..storage import Change, Row

HEADER_FILL = PatternFill("solid", fgColor="1F2937")
HEADER_FONT = Font(bold=True, color="FFFFFF")
DROP = (PatternFill("solid", fgColor="FDE2E1"), Font(color="B42318", bold=True))
RISE = (PatternFill("solid", fgColor="DCFCE7"), Font(color="15803D"))
PERCENT = '+0.0%;-0.0%;"—"'


def sort_key(r: Row):
    # Сначала подешевевшие (самые сильные сверху), потом подорожавшие, потом без изменений
    p = r.percent
    return (0 if p is not None and p < 0 else 1 if p is not None and p > 0 else 2, p or 0, r.source, r.title)


def _style_sheet(ws, widths: list[int]) -> None:
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
    for cell in ws[1]:
        cell.fill, cell.font = HEADER_FILL, HEADER_FONT
        cell.alignment = Alignment(vertical="center")
    ws.row_dimensions[1].height = 22
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions


def export(rows: list[Row], history: list[tuple[int, Change]], path: str | Path) -> Path:
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()

    ws = wb.active
    ws.title = "Цены"
    ws.append(["Конкурент", "Товар", "Цена", "Было", "Изменение", "Мин. за 30 дней", "Цена изменилась",
               "Проверено", "Ссылка"])
    for r in sorted(rows, key=sort_key):
        ws.append([
            r.source, r.title, r.price_cents / 100, r.prev_cents / 100 if r.prev_cents else None,
            r.percent / 100 if r.percent is not None else None, r.min_30d_cents / 100,
            datetime.fromtimestamp(r.changed_at), datetime.fromtimestamp(r.last_seen), r.url,
        ])
        row = ws.max_row
        for col in (3, 4, 6):
            ws.cell(row, col).number_format = excel_format(r.currency)
        ws.cell(row, 5).number_format = PERCENT
        ws.cell(row, 7).number_format = ws.cell(row, 8).number_format = "DD.MM.YYYY HH:MM"
        link = ws.cell(row, 9)
        link.hyperlink, link.value, link.style = r.url, "открыть", "Hyperlink"
    last = max(ws.max_row, 2)
    ws.conditional_formatting.add(f"E2:E{last}", CellIsRule(operator="lessThan", formula=["0"],
                                                            fill=DROP[0], font=DROP[1]))
    ws.conditional_formatting.add(f"E2:E{last}", CellIsRule(operator="greaterThan", formula=["0"],
                                                            fill=RISE[0], font=RISE[1]))
    _style_sheet(ws, [20, 60, 12, 12, 12, 16, 17, 17, 10])

    hs = wb.create_sheet("Изменения")
    hs.append(["Когда", "Конкурент", "Товар", "Было", "Стало", "Изменение", "Ссылка"])
    for ts, c in history:
        hs.append([datetime.fromtimestamp(ts), c.source, c.title, c.old_cents / 100, c.new_cents / 100,
                   c.percent / 100, c.url])
        row = hs.max_row
        hs.cell(row, 1).number_format = "DD.MM.YYYY HH:MM"
        hs.cell(row, 4).number_format = hs.cell(row, 5).number_format = excel_format(c.currency)
        hs.cell(row, 6).number_format = PERCENT
        hs.cell(row, 7).hyperlink, hs.cell(row, 7).value, hs.cell(row, 7).style = c.url, "открыть", "Hyperlink"
    last = max(hs.max_row, 2)
    hs.conditional_formatting.add(f"F2:F{last}", CellIsRule(operator="lessThan", formula=["0"],
                                                            fill=DROP[0], font=DROP[1]))
    hs.conditional_formatting.add(f"F2:F{last}", CellIsRule(operator="greaterThan", formula=["0"],
                                                            fill=RISE[0], font=RISE[1]))
    _style_sheet(hs, [17, 20, 60, 12, 12, 12, 10])

    tmp = path.with_suffix(".tmp.xlsx")
    wb.save(tmp)
    tmp.replace(path)  # атомарно: открытый в Excel старый файл не превратится в битый
    return path
